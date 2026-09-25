import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.core.config import settings

class InspectionDOCXReportGenerator:
    """
    Generates editable DOCX show-cause and compliance inspection notices for Legal Metrology officers.
    """

    @classmethod
    def generate(cls, inspection_data: dict, output_filename: str) -> str:
        filepath = os.path.join(settings.REPORTS_DIR, output_filename)
        doc = Document()

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run("GOVERNMENT OF INDIA\nDEPARTMENT OF LEGAL METROLOGY\nINSPECTION & SHOW-CAUSE NOTICE")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(15, 23, 42)

        doc.add_paragraph("Issued under Section 18 / Section 36 of the Legal Metrology Act, 2009").alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        est = inspection_data.get("establishment", {})
        doc.add_heading("1. Inspection Information", level=2)
        doc.add_paragraph(f"Inspection Notice Ref: {inspection_data.get('inspection_number')}")
        doc.add_paragraph(f"Establishment / Retailer: {est.get('name', 'N/A')}")
        doc.add_paragraph(f"Address: {est.get('address', 'N/A')}, {est.get('city', '')} {est.get('pincode', '')}")
        doc.add_paragraph(f"Date of Inspection: {inspection_data.get('scheduled_date', '')}")

        # Products
        doc.add_heading("2. Inspected Packaged Commodities & Findings", level=2)
        for p in inspection_data.get("products", []):
            p_head = doc.add_paragraph()
            p_run = p_head.add_run(f"Product: {p.get('product_name')} ({p.get('category')})")
            p_run.bold = True
            
            status_p = doc.add_paragraph(f"Status: {p.get('compliance_status')} | Readability: {p.get('readability_grade', 'N/A')} ({p.get('readability_score', 0)}%)")
            
            findings = p.get("findings", [])
            if findings:
                table = doc.add_table(rows=1, cols=4)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = "Rule / Requirement"
                hdr_cells[1].text = "Category"
                hdr_cells[2].text = "Severity"
                hdr_cells[3].text = "Finding Details"
                
                for f in findings:
                    row_cells = table.add_row().cells
                    row_cells[0].text = f.get("requirement", "")
                    row_cells[1].text = str(f.get("category", ""))
                    row_cells[2].text = str(f.get("severity", ""))
                    row_cells[3].text = f.get("reason", "")
            else:
                doc.add_paragraph("No statutory non-compliance observed on this package.")

        doc.add_heading("3. Notice & Directives", level=2)
        doc.add_paragraph(
            "Take notice that the above packaged commodities were inspected by authorized officers. "
            "Any confirmed violations are subject to penalties under Section 36 of the Legal Metrology Act, 2009. "
            "You are hereby directed to submit your explanation within 7 working days."
        )

        doc.add_paragraph("\n\n_____________________________________\nAuthorized Legal Metrology Officer\nSeal & Signature")

        doc.save(filepath)
        return filepath
